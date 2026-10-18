"""Tests for reporting (Excel, CSV, HTML)."""
import io

import pytest
from openpyxl import load_workbook

from riskreg.db import RiskDB
from riskreg.reporting import (build_csv, build_excel_report,
                               build_html_report)


@pytest.fixture()
def db(tmp_path):
    d = RiskDB(tmp_path / "r.db")
    d.create_risk(dict(title="Unpatched servers", category="Cyber Security",
                       owner="CISO", department="Security",
                       inherent_likelihood=4, inherent_impact=5,
                       residual_likelihood=2, residual_impact=4,
                       treatment="Mitigate", status="Assessed",
                       existing_controls="patch cycle",
                       identified_date="2025-01-10", next_review_date="2025-07-01"))
    d.create_risk(dict(title="Cloud dependency", category="Technology / IT Operations",
                       owner="IT Lead", department="IT",
                       inherent_likelihood=4, inherent_impact=5,
                       residual_likelihood=4, residual_impact=5,
                       treatment="Accept", status="In Treatment",
                       existing_controls="multi-region",
                       identified_date="2025-01-10", next_review_date="2025-07-01"))
    return d


def test_excel_report_loads(db):
    data = build_excel_report(db)
    wb = load_workbook(io.BytesIO(data))
    assert set(wb.sheetnames) == {"Risk Register", "Mitigation Actions",
                                  "Summary", "Heatmap data"}
    ws = wb["Risk Register"]
    assert ws.cell(row=1, column=1).value == "risk_id"
    assert ws.cell(row=2, column=1).value in {"RSK-001", "RSK-002"}
    # rating cells are filled
    fills = {str(ws.cell(row=2, column=c).fill.fgColor.rgb).lower()
             for c in range(1, ws.max_column + 1)}
    assert any(x in f for f in fills for x in ("2e9e4f", "f2c230", "d62f2f"))
    ws3 = wb["Summary"]
    assert ws3.cell(row=1, column=1).value == "Metric"


def test_csv_report(db):
    csv = build_csv(db)
    assert "risk_id" in csv and "RSK-001" in csv and "RSK-002" in csv


def test_html_report(db):
    html = build_html_report(db)
    assert "Executive Report" in html
    assert "Inherent Risk Heatmap" in html
    assert "Residual Risk Heatmap" in html
    assert "Top 10 Risks" in html
    assert "Appetite Breaches" in html
    assert "Methodology" in html
    assert "Unpatched servers" in html
