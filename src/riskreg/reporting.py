"""Excel, CSV and HTML report generation."""
from __future__ import annotations

from datetime import date
from html import escape
from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from .analytics import (category_breakdown, heatmap_matrix,
                        outside_appetite_table, overdue_reviews)
from .config import Config, load_config
from .scoring import risk_reduction_pct

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
BAND_COLORS = {
    "Low": "2e9e4f",
    "Medium": "f2c230",
    "High": "f08a24",
    "Critical": "d62f2f",
}
BAND_FILLS = {name: PatternFill("solid", fgColor=hexv)
              for name, hexv in BAND_COLORS.items()}
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _style_header(ws, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN


def _autowidth(ws, df: pd.DataFrame, start_row: int) -> None:
    for col_idx, col_name in enumerate(df.columns, start=1):
        max_len = len(str(col_name))
        for val in df[col_name].head(40):
            max_len = max(max_len, len(str(val)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(
            max(max_len + 2, 10), 52)


def _write_df(ws, df: pd.DataFrame, start_row: int = 1, header: bool = True,
              rating_cols: tuple[str, ...] = ()) -> None:
    if header:
        for col_idx, col_name in enumerate(df.columns, start=1):
            ws.cell(row=start_row, column=col_idx, value=str(col_name))
        _style_header(ws, start_row, len(df.columns))
        data_row = start_row + 1
    else:
        data_row = start_row
    for r, (_, row) in enumerate(df.iterrows()):
        for c, col_name in enumerate(df.columns, start=1):
            cell = ws.cell(row=data_row + r, column=c, value=row[col_name])
            cell.border = THIN
            cell.alignment = LEFT if c == 1 else CENTER
            if col_name in rating_cols and str(row[col_name]) in BAND_FILLS:
                cell.fill = BAND_FILLS[str(row[col_name])]
                cell.font = Font(bold=True, color="FFFFFF", size=9)
    _autowidth(ws, df, start_row)


def build_excel_report(db) -> bytes:
    """Build a formatted Excel workbook (bytes) with four sheets."""
    wb = Workbook()
    risks = db.list_risks()
    actions = db.list_actions()

    ws = wb.active
    ws.title = "Risk Register"
    risk_cols = ["risk_id", "title", "category", "owner", "department",
                 "inherent_score", "inherent_rating", "residual_score",
                 "residual_rating", "treatment", "status", "next_review_date"]
    _write_df(ws, risks[risk_cols], rating_cols=("inherent_rating",
                                                 "residual_rating"))

    ws2 = wb.create_sheet("Mitigation Actions")
    act_cols = ["action_id", "risk_id", "risk_title", "description",
                "action_owner", "due_date", "status", "completed_date"]
    _write_df(ws2, actions[act_cols], rating_cols=("status",))

    ws3 = wb.create_sheet("Summary")
    config = db.config
    summary_rows = [
        ("Total open risks", int((risks["is_deleted"] == 0).sum())),
        ("Critical risks", int((risks["residual_rating"] == "Critical").sum())),
        ("High risks", int((risks["residual_rating"] == "High").sum())),
        ("Outside appetite",
         int(outside_appetite_table(risks, config).shape[0])),
        ("Overdue reviews", int(overdue_reviews(risks, date.today()).shape[0])),
        ("Overdue actions",
         int((actions["status"] == "Overdue").sum())),
        ("Average risk reduction %",
         round(float(risks.apply(
             lambda r: risk_reduction_pct(int(r["inherent_score"]),
                                          int(r["residual_score"])), axis=1)
                 .mean()), 1) if len(risks) else 0.0),
    ]
    ws3.append(["Metric", "Value"])
    _style_header(ws3, 1, 2)
    for i, (k, v) in enumerate(summary_rows, start=2):
        ws3.cell(row=i, column=1, value=k).border = THIN
        ws3.cell(row=i, column=1).alignment = LEFT
        c = ws3.cell(row=i, column=2, value=v); c.border = THIN
        c.alignment = CENTER
    ws3.column_dimensions["A"].width = 30
    ws3.column_dimensions["B"].width = 14

    ws4 = wb.create_sheet("Heatmap data")
    ws4.append(["Inherent heatmap (counts by likelihood x impact)"])
    counts, _, matrix = heatmap_matrix(risks, config, "inherent")
    _write_heatmap_sheet(ws4, counts, start_row=3)
    start = 3 + 6 + 2
    ws4.cell(row=start, column=1,
             value="Residual heatmap (counts by likelihood x impact)")
    counts2, _, _ = heatmap_matrix(risks, config, "residual")
    _write_heatmap_sheet(ws4, counts2, start_row=start + 2)
    _autowidth(ws4, pd.DataFrame(), 1)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_heatmap_sheet(ws, counts, start_row: int) -> None:
    ws.cell(row=start_row, column=1, value="Likelihood \\ Impact")
    for i in range(1, 6):
        ws.cell(row=start_row, column=1 + i, value=f"Impact {i}")
    for l in range(1, 6):
        ws.cell(row=start_row + l, column=1, value=f"Likelihood {l}")
        for i in range(1, 6):
            ws.cell(row=start_row + l, column=1 + i, value=int(counts[l][i]))

def _heatmap_html(risks: pd.DataFrame, config: Config, score_type: str,
                  title: str) -> str:
    counts, details, matrix = heatmap_matrix(risks, config, score_type)
    band_colors = {str(b["name"]): str(b["color"]) for b in
                   config["rating_bands"]}
    order = ["Low", "Medium", "High", "Critical"]
    cells = []
    for l in range(1, 6):
        for i in range(1, 6):
            items = details.get((l, i), [])
            rating_idx = matrix[l - 1][i - 1]
            rating = order[rating_idx] if rating_idx is not None else ""
            color = band_colors.get(rating, "#ffffff") if rating else "#f5f5f5"
            tooltip = escape(title)
            if items:
                tooltip += "<ul>" + "".join(
                    f"<li>{escape(d['risk_id'])}: {escape(d['title'])} "
                    f"(score {d['score']})</li>" for d in items) + "</ul>"
            cells.append(
                f"<td class='cell' style='background:{color}' "
                f"title='{tooltip}'>"
                f"<div class='count'>{len(items)}</div>"
                f"<div class='rating'>{rating}</div></td>")
    rows_html = "".join(
        f"<tr>{''.join(cells[r * 5:(r + 1) * 5])}</tr>" for r in range(5))
    def _label(scale, n):
        entry = scale.get(n) or scale.get(str(n))
        return f"{n} {entry['label']}"
    impact_labels = [_label(config["impact_scale"], i) for i in range(1, 6)]
    likelihood_labels = [_label(config["likelihood_scale"], l)
                         for l in range(1, 6)]
    header_cells = "".join(f"<th>{escape(lb)}</th>" for lb in impact_labels)
    y_cells = "".join(
        f"<th class='y'>{escape(lb)}</th>" for lb in likelihood_labels)
    return f"""
    <h3>{escape(title)}</h3>
    <table class='heatmap'>
      <tr><th class='corner'>Likelihood \\\\ Impact</th>{header_cells}</tr>
      {rows_html}
    </table>
    <div class='legend'>
      {''.join(f"<span style='background:{band_colors[n]}'>{n}</span>" for n in order)}
    </div>
    """


def build_html_report(db) -> str:
    """Return a self-contained HTML executive report."""
    config = db.config
    risks = db.list_risks()
    actions = db.list_actions()
    today = date.today().isoformat()
    summary_rows = [
        ("Total open risks", int((risks["is_deleted"] == 0).sum())),
        ("Critical risks", int((risks["residual_rating"] == "Critical").sum())),
        ("High risks", int((risks["residual_rating"] == "High").sum())),
        ("Outside appetite", int(outside_appetite_table(risks, config).shape[0])),
        ("Overdue reviews", int(overdue_reviews(risks, date.today()).shape[0])),
        ("Overdue actions", int((actions["status"] == "Overdue").sum())),
    ]
    avg_red = (float(risks.apply(
        lambda r: risk_reduction_pct(int(r["inherent_score"]),
                                     int(r["residual_score"])), axis=1).mean())
               if len(risks) else 0.0)
    summary_rows.append(("Average risk reduction %", round(avg_red, 1)))

    top10 = (risks.sort_values(["residual_score", "risk_id"],
                               ascending=[False, True]).head(10))
    breaches = outside_appetite_table(risks, config)
    overdue_acts = actions[actions["status"] == "Overdue"]

    def df_table(df, cols, rating_cols=()):
        if df.empty:
            return "<p>None.</p>"
        thead = "".join(f"<th>{escape(c)}</th>" for c in cols)
        body = ""
        for _, row in df.iterrows():
            tds = []
            for c in cols:
                v = str(row[c])
                style = ""
                if c in rating_cols and v in BAND_COLORS:
                    style = f"style='background:#{BAND_COLORS[v]};color:#fff;font-weight:bold'"
                tds.append(f"<td {style}>{escape(v)}</td>")
            body += "<tr>" + "".join(tds) + "</tr>"
        return f"<table><thead><tr>{thead}</tr></thead><tbody>{body}</tbody></table>"

    css = """
    body{font-family:Segoe UI,Helvetica,Arial,sans-serif;color:#222;margin:24px;}
    h1{color:#1F3864;border-bottom:2px solid #1F3864;padding-bottom:8px;}
    h3{color:#1F3864;margin-top:20px;}
    table{border-collapse:collapse;width:100%;margin:10px 0;font-size:12px;}
    th,td{border:1px solid #ccc;padding:5px 8px;text-align:left;}
    th{background:#1F3864;color:#fff;}
    .kpis{display:flex;gap:12px;margin:16px 0;flex-wrap:wrap;}
    .kpi{background:#f4f6fb;border:1px solid #d7deed;border-radius:8px;
         padding:12px 16px;min-width:150px;}
    .kpi .v{font-size:24px;font-weight:bold;color:#1F3864;}
    .kpi .l{font-size:12px;color:#555;margin-top:2px;}
    .heatmap{border-collapse:collapse;margin:12px auto;font-size:11px;
              text-align:center;}
    .heatmap td{width:46px;height:46px;border:1px solid #fff;}
    .heatmap th{width:46px;height:46px;}
    .heatmap .count{font-weight:bold;font-size:13px;}
    .heatmap .rating{font-size:9px;color:#333;margin-top:2px;}
    .legend{margin:8px 0;font-size:12px;}
    .legend span{display:inline-block;width:20px;height:20px;margin-right:6px;
                 border:1px solid #999;vertical-align:middle;}
    .methodology{background:#f4f6fb;padding:14px;border-radius:8px;}
    @media print{body{margin:0;} .kpis{page-break-inside:avoid;}}
    """
    return f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8">
<title>Risk Register Executive Report</title><style>{css}</style></head>
<body>
<h1>Automated Risk Register &amp; Heatmap Dashboard — Executive Report</h1>
<p><strong>Generated:</strong> {escape(today)} &nbsp;|&nbsp;
   <strong>Appetite (max residual score):</strong>
   {config['risk_appetite']['max_residual_score']}</p>
<div class="kpis">
{''.join(f"<div class='kpi'><div class='v'>{v}</div><div class='l'>{escape(k)}</div></div>" for k, v in summary_rows)}
</div>
<h3>Risk Heatmaps</h3>
{_heatmap_html(risks, config, "inherent", "Inherent Risk Heatmap")}
{_heatmap_html(risks, config, "residual", "Residual Risk Heatmap")}
<h3>Top 10 Risks (by residual score)</h3>
{df_table(top10, ["risk_id", "title", "category", "owner",
                  "residual_score", "residual_rating"],
          rating_cols=("residual_rating",))}
<h3>Appetite Breaches (residual &gt; appetite)</h3>
{df_table(breaches, ["risk_id", "title", "category", "owner",
                     "residual_score", "residual_rating", "gap"],
          rating_cols=("residual_rating",))}
<h3>Overdue Mitigation Actions</h3>
{df_table(overdue_acts, ["risk_id", "description", "action_owner",
                         "due_date", "status"],
          rating_cols=("status",))}
<div class="methodology">
<h3>Methodology</h3>
<p>Scores are computed as Likelihood × Impact on a 1–5 scale and mapped to
rating bands: <strong>Low (1–4)</strong>, <strong>Medium (5–9)</strong>,
<strong>High (10–15)</strong>, <strong>Critical (16–25)</strong>.
Residual risk reflects the effect of existing controls and treatment and must
not exceed inherent risk. A risk is outside appetite when its residual score
exceeds the configured threshold (default
{config['risk_appetite']['max_residual_score']}; per-category overrides may
apply). Risk reduction % = (inherent − residual) / inherent × 100.</p>
<h4>Likelihood scale</h4>
<ul>{''.join(f"<li><b>{l} {(config['likelihood_scale'].get(l) or config['likelihood_scale'].get(str(l)))['label']}</b> — {escape((config['likelihood_scale'].get(l) or config['likelihood_scale'].get(str(l)))['description'])}</li>" for l in range(1, 6))}</ul>
<h4>Impact scale</h4>
<ul>{''.join(f"<li><b>{i} {(config['impact_scale'].get(i) or config['impact_scale'].get(str(i)))['label']}</b></li>" for i in range(1, 6))}</ul>
<p><strong>Treatment options:</strong>
{escape(", ".join(config['treatment_options']))}</p>
<p><strong>Status values:</strong>
{escape(", ".join(config['status_values']))}</p>
</div>
</body></html>"""


def build_csv(db) -> str:
    """Return the risk register as a CSV string."""
    return db.list_risks().to_csv(index=False)


def main() -> None:
    """CLI: write Excel, CSV and HTML reports to reports/."""
    import logging
    from pathlib import Path
    logging.basicConfig(level=logging.INFO)
    out = Path("reports")
    out.mkdir(exist_ok=True)
    from .db import RiskDB
    database = RiskDB(config=load_config())
    (out / "risk_register.xlsx").write_bytes(build_excel_report(database))
    (out / "risk_register.csv").write_text(build_csv(database), encoding="utf-8")
    (out / "risk_report.html").write_text(build_html_report(database),
                                          encoding="utf-8")
    print(f"Reports written to {out.resolve()}")


if __name__ == "__main__":
    main()

